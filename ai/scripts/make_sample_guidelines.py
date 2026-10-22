"""Regenerates the downloadable sample guideline files judges can upload (needs requirements-dev.txt)."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from app.sample_brand import SAMPLE_GUIDELINE_TEXT  # noqa: E402

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "web", "public", "sample")
os.makedirs(OUT, exist_ok=True)

from docx import Document  # noqa: E402
doc = Document()
for i, line in enumerate(SAMPLE_GUIDELINE_TEXT.splitlines()):
    if line.strip():
        doc.add_heading(line, 0) if i == 0 else doc.add_paragraph(line)
doc.save(os.path.join(OUT, "lumina-brand-guidelines.docx"))

from fpdf import FPDF  # noqa: E402
pdf = FPDF(); pdf.add_page(); pdf.set_font("Helvetica", size=11)
for line in SAMPLE_GUIDELINE_TEXT.splitlines():
    pdf.multi_cell(0, 6, line.replace("’", "'").replace("•", "-") or " ", new_x="LMARGIN", new_y="NEXT")
pdf.output(os.path.join(OUT, "lumina-brand-guidelines.pdf"))
open(os.path.join(OUT, "lumina-brand-guidelines.md"), "w").write(SAMPLE_GUIDELINE_TEXT)
print("wrote", sorted(os.listdir(OUT)))
